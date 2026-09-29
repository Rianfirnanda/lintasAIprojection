"""Lapisan pengambilan & integrasi: membaca berkas CSV/Excel dari data/masuk menjadi observasi terstandar.

Setiap berkas dicatat metadatanya (checksum SHA-256, jumlah baris diterima/ditolak) sebagai jejak audit.
Baris yang gagal validasi skema ditolak dengan alasan yang jelas; baris yang lolos diteruskan ke quality gate.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from .konfigurasi import Konfigurasi

# Konversi satuan masukan -> satuan dasar varian (faktor pengali harga).
# Contoh: harga per ons x 10 = harga per kg.
KONVERSI_SATUAN = {
    ("ons", "kg"): 10.0,
    ("gram", "kg"): 1000.0,
    ("g", "kg"): 1000.0,
    ("ml", "liter"): 1000.0,
    ("l", "liter"): 1.0,
    ("lt", "liter"): 1.0,
    ("ltr", "liter"): 1.0,
    ("kilogram", "kg"): 1.0,
}

ALIAS_HARGA = {
    "tanggal": ["tanggal", "tgl", "date", "tanggal_pencatatan"],
    "kode_pasar": ["kode_pasar", "pasar", "market"],
    "kode_varian": ["kode_varian", "varian", "kode_komoditas_varian", "kode_barang"],
    "harga": ["harga", "harga_rp", "price", "harga_eceran"],
    "satuan": ["satuan", "unit"],
    "kode_sumber": ["kode_sumber", "sumber"],
    "petugas": ["petugas", "kode_petugas", "pencacah"],
    "responden": ["responden", "kode_responden"],
    "catatan": ["catatan", "keterangan", "notes"],
    "id_klien": ["id_klien", "uuid", "client_id"],
    "waktu_input": ["waktu_input", "waktu_isi", "timestamp"],
}

ALIAS_KONTEKS = {
    "tanggal": ["tanggal", "tgl", "date"],
    "kode_wilayah": ["kode_wilayah", "wilayah"],
    "indikator": ["indikator", "variabel", "indicator"],
    "nilai": ["nilai", "value"],
    "satuan": ["satuan", "unit"],
    "kode_varian": ["kode_varian", "varian"],
    "kode_sumber": ["kode_sumber", "sumber"],
}

WAJIB_HARGA = ("tanggal", "kode_pasar", "kode_varian", "harga")
WAJIB_KONTEKS = ("tanggal", "kode_wilayah", "indikator", "nilai")


@dataclass
class Observasi:
    id: str
    tanggal: date
    kode_pasar: str
    kode_wilayah: str
    kode_varian: str
    kode_sumber: str
    harga: float  # dalam satuan dasar varian
    harga_asli: float
    satuan_asli: str
    petugas: str
    responden: str
    id_klien: str
    waktu_input: datetime | None
    berkas: str
    baris: int
    catatan: str = ""
    status: str = "baru"
    tanda: list[str] = field(default_factory=list)


@dataclass
class ObservasiKonteks:
    tanggal: date
    kode_wilayah: str
    indikator: str
    nilai: float
    satuan: str
    kode_varian: str
    kode_sumber: str
    berkas: str


@dataclass
class Penolakan:
    berkas: str
    baris: int
    alasan: str
    data: dict


@dataclass
class Batch:
    berkas: str
    jenis: str
    sha256: str
    ukuran_byte: int
    jumlah_baris: int = 0
    diterima: int = 0
    ditolak: int = 0
    galat_berkas: str = ""


@dataclass
class HasilMasukan:
    observasi: list[Observasi] = field(default_factory=list)
    konteks: list[ObservasiKonteks] = field(default_factory=list)
    penolakan: list[Penolakan] = field(default_factory=list)
    batch: list[Batch] = field(default_factory=list)


def normalisasi_kolom(nama: str) -> str:
    nama = str(nama or "").strip().lower()
    nama = re.sub(r"[^a-z0-9]+", "_", nama)
    return nama.strip("_")


def _peta_alias(kolom: list[str], alias: dict[str, list[str]]) -> dict[str, str]:
    """Kembalikan pemetaan nama_kanonik -> nama kolom di berkas."""
    peta = {}
    for kanonik, pilihan in alias.items():
        for p in pilihan:
            if p in kolom:
                peta[kanonik] = p
                break
    return peta


def baca_tabel(path: Path) -> tuple[list[str], list[dict]]:
    """Baca CSV (pemisah , atau ;) atau XLSX (sheet pertama). Header dinormalisasi."""
    akhiran = path.suffix.lower()
    if akhiran in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        baris_iter = ws.iter_rows(values_only=True)
        try:
            header = [normalisasi_kolom(h) for h in next(baris_iter)]
        except StopIteration:
            return [], []
        hasil = []
        for nilai in baris_iter:
            if nilai is None or all(v is None or str(v).strip() == "" for v in nilai):
                continue
            hasil.append({header[i]: nilai[i] for i in range(min(len(header), len(nilai))) if header[i]})
        wb.close()
        return header, hasil

    teks = path.read_bytes().decode("utf-8-sig")
    contoh = teks[:4096]
    pemisah = ";" if contoh.count(";") > contoh.count(",") else ","
    pembaca = csv.reader(io.StringIO(teks), delimiter=pemisah)
    try:
        header = [normalisasi_kolom(h) for h in next(pembaca)]
    except StopIteration:
        return [], []
    hasil = []
    for nilai in pembaca:
        if not nilai or all(not str(v).strip() for v in nilai):
            continue
        hasil.append({header[i]: nilai[i].strip() for i in range(min(len(header), len(nilai))) if header[i]})
    return header, hasil


def parse_tanggal(nilai) -> date:
    if isinstance(nilai, datetime):
        return nilai.date()
    if isinstance(nilai, date):
        return nilai
    if isinstance(nilai, (int, float)):
        # nomor seri tanggal Excel
        return date(1899, 12, 30) + timedelta(days=int(nilai))
    teks = str(nilai).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(teks, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"format tanggal tidak dikenali: {teks!r}")


def parse_waktu(nilai) -> datetime | None:
    if nilai in (None, ""):
        return None
    if isinstance(nilai, datetime):
        return nilai
    teks = str(nilai).strip().replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(teks)
    except ValueError:
        return None


_RIBUAN_TITIK = re.compile(r"^\d{1,3}(\.\d{3})+$")
_RIBUAN_KOMA = re.compile(r"^\d{1,3}(,\d{3})+$")


def parse_harga(nilai) -> float:
    """Terima angka biasa atau format Indonesia: 'Rp 45.000', '45.000', '12.500,50'."""
    if isinstance(nilai, (int, float)):
        angka = float(nilai)
    else:
        teks = str(nilai).strip().lower().replace("rp", "").replace(" ", "")
        if not teks:
            raise ValueError("harga kosong")
        if "." in teks and "," in teks:
            teks = teks.replace(".", "").replace(",", ".")
        elif _RIBUAN_TITIK.match(teks):
            teks = teks.replace(".", "")
        elif _RIBUAN_KOMA.match(teks):
            teks = teks.replace(",", "")
        elif "," in teks:
            teks = teks.replace(",", ".")
        angka = float(teks)
    if angka <= 0:
        raise ValueError("harga harus lebih dari 0")
    return angka


def faktor_konversi(satuan_asli: str, satuan_dasar: str) -> float:
    a = satuan_asli.strip().lower()
    d = satuan_dasar.strip().lower()
    if not a or a == d:
        return 1.0
    if (a, d) in KONVERSI_SATUAN:
        return KONVERSI_SATUAN[(a, d)]
    raise ValueError(f"satuan '{satuan_asli}' tidak dapat dikonversi ke '{satuan_dasar}'")


def id_observasi(*bagian) -> str:
    return hashlib.sha1("|".join(str(b) for b in bagian).encode()).hexdigest()[:12]


def sha256_berkas(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for blok in iter(lambda: f.read(65536), b""):
            h.update(blok)
    return h.hexdigest()


def _cek_privasi(header: list[str], konf: Konfigurasi) -> str:
    terlarang = set(konf.pengaturan["privasi"]["kolom_terlarang"])
    kena = sorted(set(header) & terlarang)
    if kena:
        return (
            "berkas ditolak: memuat kolom data pribadi "
            f"({', '.join(kena)}). Hapus kolom tersebut; gunakan kode responden anonim."
        )
    return ""


def _nilai(baris: dict, peta: dict, kunci: str, default=""):
    kolom = peta.get(kunci)
    if kolom is None:
        return default
    v = baris.get(kolom)
    return default if v is None else v


def baca_berkas_harga(path: Path, relatif: str, konf: Konfigurasi, hasil: HasilMasukan) -> None:
    batch = Batch(berkas=relatif, jenis="harga", sha256=sha256_berkas(path), ukuran_byte=path.stat().st_size)
    hasil.batch.append(batch)
    try:
        header, daftar = baca_tabel(path)
    except Exception as e:  # berkas rusak/tidak terbaca
        batch.galat_berkas = f"berkas tidak dapat dibaca: {e}"
        return
    batch.jumlah_baris = len(daftar)
    galat = _cek_privasi(header, konf)
    peta = _peta_alias(header, ALIAS_HARGA)
    kurang = [k for k in WAJIB_HARGA if k not in peta]
    if not galat and kurang:
        galat = f"kolom wajib tidak ditemukan: {', '.join(kurang)}"
    if galat:
        batch.galat_berkas = galat
        batch.ditolak = len(daftar)
        return

    for i, baris in enumerate(daftar, start=2):  # baris 1 = header
        data_asli = {k: (v.isoformat() if isinstance(v, (date, datetime)) else v) for k, v in baris.items()}
        try:
            mentah_tanggal = _nilai(baris, peta, "tanggal")
            if mentah_tanggal in ("", None):
                raise ValueError("tanggal kosong")
            tgl = parse_tanggal(mentah_tanggal)
            if tgl > konf.hari_ini:
                raise ValueError(f"tanggal {tgl} melewati hari ini")
            kode_pasar = str(_nilai(baris, peta, "kode_pasar")).strip().upper()
            if not kode_pasar:
                raise ValueError("kode_pasar kosong")
            pasar = konf.pasar.get(kode_pasar)
            if pasar is None:
                raise ValueError(f"kode_pasar tidak dikenal: {kode_pasar}")
            kode_varian = str(_nilai(baris, peta, "kode_varian")).strip().upper()
            if not kode_varian:
                raise ValueError("kode_varian kosong")
            varian = konf.varian.get(kode_varian)
            if varian is None:
                raise ValueError(f"kode_varian tidak dikenal: {kode_varian}")
            mentah_harga = _nilai(baris, peta, "harga")
            if mentah_harga in ("", None):
                raise ValueError("harga kosong")
            harga_asli = parse_harga(mentah_harga)
            satuan_asli = str(_nilai(baris, peta, "satuan") or varian.satuan).strip().lower()
            harga = harga_asli * faktor_konversi(satuan_asli, varian.satuan)
            kode_sumber = str(_nilai(baris, peta, "kode_sumber") or "PSR-ENUM").strip().upper()
            if kode_sumber not in konf.sumber:
                raise ValueError(f"kode_sumber tidak dikenal: {kode_sumber}")
        except ValueError as e:
            hasil.penolakan.append(Penolakan(berkas=relatif, baris=i, alasan=str(e), data=data_asli))
            batch.ditolak += 1
            continue

        petugas = str(_nilai(baris, peta, "petugas")).strip()
        responden = str(_nilai(baris, peta, "responden")).strip()
        id_klien = str(_nilai(baris, peta, "id_klien")).strip()
        hasil.observasi.append(
            Observasi(
                id=id_observasi(tgl, kode_pasar, kode_varian, kode_sumber, responden, harga_asli, satuan_asli, id_klien),
                tanggal=tgl, kode_pasar=kode_pasar, kode_wilayah=pasar.kode_wilayah,
                kode_varian=kode_varian, kode_sumber=kode_sumber, harga=round(harga, 2),
                harga_asli=harga_asli, satuan_asli=satuan_asli, petugas=petugas,
                responden=responden, id_klien=id_klien,
                waktu_input=parse_waktu(_nilai(baris, peta, "waktu_input", None)),
                berkas=relatif, baris=i, catatan=str(_nilai(baris, peta, "catatan")),
            )
        )
        batch.diterima += 1


def baca_berkas_konteks(path: Path, relatif: str, konf: Konfigurasi, hasil: HasilMasukan) -> None:
    batch = Batch(berkas=relatif, jenis="konteks", sha256=sha256_berkas(path), ukuran_byte=path.stat().st_size)
    hasil.batch.append(batch)
    try:
        header, daftar = baca_tabel(path)
    except Exception as e:
        batch.galat_berkas = f"berkas tidak dapat dibaca: {e}"
        return
    batch.jumlah_baris = len(daftar)
    galat = _cek_privasi(header, konf)
    peta = _peta_alias(header, ALIAS_KONTEKS)
    kurang = [k for k in WAJIB_KONTEKS if k not in peta]
    if not galat and kurang:
        galat = f"kolom wajib tidak ditemukan: {', '.join(kurang)}"
    if galat:
        batch.galat_berkas = galat
        batch.ditolak = len(daftar)
        return
    for i, baris in enumerate(daftar, start=2):
        try:
            tgl = parse_tanggal(_nilai(baris, peta, "tanggal"))
            wil = str(_nilai(baris, peta, "kode_wilayah")).strip()
            if wil not in konf.wilayah:
                raise ValueError(f"kode_wilayah tidak dikenal: {wil}")
            indikator = normalisasi_kolom(_nilai(baris, peta, "indikator"))
            if not indikator:
                raise ValueError("indikator kosong")
            mentah = _nilai(baris, peta, "nilai")
            if mentah in ("", None):
                raise ValueError("nilai kosong")
            nilai = float(str(mentah).replace(",", ".")) if not isinstance(mentah, (int, float)) else float(mentah)
            kode_varian = str(_nilai(baris, peta, "kode_varian")).strip().upper()
            if kode_varian and kode_varian not in konf.varian:
                raise ValueError(f"kode_varian tidak dikenal: {kode_varian}")
        except ValueError as e:
            hasil.penolakan.append(Penolakan(berkas=relatif, baris=i, alasan=str(e), data=dict(baris)))
            batch.ditolak += 1
            continue
        hasil.konteks.append(
            ObservasiKonteks(
                tanggal=tgl, kode_wilayah=wil, indikator=indikator, nilai=nilai,
                satuan=str(_nilai(baris, peta, "satuan")), kode_varian=kode_varian,
                kode_sumber=str(_nilai(baris, peta, "kode_sumber")).strip().upper(),
                berkas=relatif,
            )
        )
        batch.diterima += 1


def _daftar_berkas(folder: Path) -> list[Path]:
    if not folder.exists():
        return []
    return sorted(
        p for p in folder.rglob("*")
        if p.is_file() and p.suffix.lower() in (".csv", ".xlsx", ".xlsm") and not p.name.startswith(("~$", "."))
    )


def baca_semua(folder_masuk: Path, konf: Konfigurasi, akar_relatif: Path | None = None) -> HasilMasukan:
    """Baca seluruh berkas di <folder_masuk>/harga dan <folder_masuk>/konteks."""
    hasil = HasilMasukan()
    akar_relatif = akar_relatif or folder_masuk.parent.parent
    for p in _daftar_berkas(folder_masuk / "harga"):
        baca_berkas_harga(p, _relatif(p, akar_relatif), konf, hasil)
    for p in _daftar_berkas(folder_masuk / "konteks"):
        baca_berkas_konteks(p, _relatif(p, akar_relatif), konf, hasil)
    return hasil


def _relatif(p: Path, akar: Path) -> str:
    try:
        return p.resolve().relative_to(akar.resolve()).as_posix()
    except ValueError:
        return p.as_posix()


def ada_data_harga(folder_masuk: Path) -> bool:
    return bool(_daftar_berkas(folder_masuk / "harga"))
